# //// Neoffice — added file (no upstream equivalent).
"""A table of products as a file: CSV and JSON for another system, Excel for a person.

The three read the same lines (`rows.build_rows`), and the CSV and the JSON the same keys in the same order:
the machine names of `rows.columns_for`, English and stable whatever the language of the reseller, because a
script or a shop's importer matches on them. Only Excel is worded for a person (translated headers, "In
stock", "Yes").

- **Figures stay figures.** A price is a number in JSON and in Excel (two decimals shown, more kept when the
  price has more), and a dot-decimal text in CSV.
- **A text that starts like a formula stays text in Excel.** openpyxl turns a string that starts with "="
  into a formula; every text cell is set to a string, so a product named "=1+1" is a product. The CSV is left
  as it is: it is for another system, which matches on what is written, and an apostrophe added to a name
  would be imported with it.
- **The CSV has no byte-order mark** and ends its lines with CRLF (RFC 4180); to read a file in a
  spreadsheet, the page offers Excel.
"""

import csv
import io
import json

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRICE_KEYS = ("price", "purchase_price")
MAX_WIDTH = 60
SAMPLE = 200  # lines read to size a column


def amount_text(value):
	"""A price as text, a dot, at least two decimals and as many more as it has (up to four): 80.00, 0.125."""
	whole, _dot, decimals = f"{float(value):.4f}".partition(".")
	return f"{whole}.{decimals.rstrip('0').ljust(2, '0')}"


def attributes_text(attributes):
	return "; ".join(f"{name}: {value}" for name, value in (attributes or {}).items())


def _csv_cell(key, value):
	if value is None:
		return ""
	if isinstance(value, bool):
		return "true" if value else "false"
	if isinstance(value, (int, float)):
		return amount_text(value)
	if key == "attributes":
		return attributes_text(value)
	if key == "more_images":
		return ",".join(value or [])
	return str(value)


def csv_bytes(data):
	"""The CSV: a header of machine names, one record per line."""
	out = io.StringIO(newline="")
	writer = csv.writer(out, lineterminator="\r\n")
	writer.writerow(data.columns)
	for line in data.rows:
		writer.writerow([_csv_cell(key, line.get(key)) for key in data.columns])
	return out.getvalue().encode("utf-8")


def json_bytes(data, generated_on=None):
	"""The JSON: the file's own facts, then one object per line with the keys of the CSV's columns."""
	products = []
	for line in data.rows:
		product = {}
		for key in data.columns:
			value = line.get(key)
			if key in PRICE_KEYS and value is not None:
				value = round(float(value), 4)
			product[key] = value
		products.append(product)
	document = {
		"generated_on": generated_on,
		"currency": data.currency if data.prices != "none" else None,
		"prices_as_of": data.prices_as_of,
		"count": len(products),
		"products": products,
	}
	return json.dumps(document, ensure_ascii=False, indent=1).encode("utf-8")


def _sheet_text(text):
	"""A control character (a vertical tab pasted from Word) makes openpyxl refuse the whole workbook; it prints
	nothing, so it goes."""
	return ILLEGAL_CHARACTERS_RE.sub("", text)


def _excel_value(key, value, labels):
	if value is None:
		return None
	if isinstance(value, bool):
		return labels["yes"] if value else labels["no"]
	if isinstance(value, (int, float)):
		return float(value)
	if key == "attributes":
		return _sheet_text(attributes_text(value))
	if key == "more_images":
		return _sheet_text("\n".join(value or []))
	if key == "availability":
		return labels["availability"].get(value, value)
	return _sheet_text(str(value))


def _as_text(cells):
	"""Every string cell is a string for Excel, whatever it starts with."""
	for cell in cells:
		if isinstance(cell.value, str):
			cell.data_type = "s"


def xlsx_bytes(data, labels):
	"""The workbook: a sheet of products and a sheet saying what the file is (`labels`: see `rows.labels`)."""
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = labels["products"]
	sheet.append([labels["columns"].get(key, key) for key in data.columns])
	for cell in sheet[1]:
		cell.font = Font(bold=True)
		cell.fill = PatternFill("solid", fgColor="EFE9DD")
		cell.alignment = Alignment(vertical="center")
	for line in data.rows:
		sheet.append([_excel_value(key, line.get(key), labels) for key in data.columns])
		_as_text(sheet[sheet.max_row])
	for index, key in enumerate(data.columns, start=1):
		letter = get_column_letter(index)
		if key in PRICE_KEYS:
			for cell in sheet[letter][1:]:
				cell.number_format = "0.00##"
		longest = max(
			[len(str(sheet.cell(row=1, column=index).value or ""))]
			+ [len(str(cell.value or "").split("\n")[0]) for cell in sheet[letter][1 : SAMPLE + 1]]
		)
		sheet.column_dimensions[letter].width = min(longest + 2, MAX_WIDTH)
		if key in ("description", "more_images"):
			for cell in sheet[letter][1:]:
				cell.alignment = Alignment(wrap_text=True, vertical="top")
	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = sheet.dimensions

	facts = workbook.create_sheet(labels["information"])
	for label, value in labels["facts"].items():
		facts.append([label, _sheet_text(value) if isinstance(value, str) else value])
		facts.cell(row=facts.max_row, column=1).font = Font(bold=True)
		_as_text(facts[facts.max_row])
	facts.column_dimensions["A"].width = 28
	facts.column_dimensions["B"].width = 60

	out = io.BytesIO()
	workbook.save(out)
	return out.getvalue()
