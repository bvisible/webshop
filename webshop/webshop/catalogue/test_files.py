# //// Neoffice — added file (no upstream equivalent).
"""The reseller's workbooks and CSV files: asked on the catalogues page, made by the same job as a PDF.

Same two resellers and the same shop as the catalogues' tests. What is proved is what differs from a PDF: the
format is asked for and kept, the price modes a file may have (their own price, for themselves) are refused to
a PDF, a file needs no Chromium and no brand when it carries no selling price, the job writes a file that opens,
and the download says what it is.
"""

import csv
import io
import json
from unittest import mock

import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import load_workbook

from webshop.webshop.catalogue import access, catalogue_api, jobs, rows
from webshop.webshop.catalogue import api as brand_api
from webshop.webshop.catalogue import fixtures as shop
from webshop.webshop.catalogue.refusals import catalogue_message
from webshop.webshop.doctype.reseller_catalogue.reseller_catalogue import FAILED, READY

DOCTYPE = "Reseller Catalogue"
PICKED = {"item_groups": [shop.PARENT_GROUP, shop.OTHER_GROUP]}


class TestFiles(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		shop.make()

	@classmethod
	def tearDownClass(cls):
		shop.purge()
		super().tearDownClass()

	def setUp(self):
		self.settings = frappe.get_doc("Webshop Settings")
		self.settings.enable_reseller_catalogue = 1
		self.settings.reseller_catalogue_groups = [frappe._dict(allowed_customer_group=shop.RESELLER_GROUP)]
		self.settings.catalogue_credit_text = ""
		patcher = mock.patch.object(access, "shop_settings", return_value=self.settings)
		patcher.start()
		self.addCleanup(patcher.stop)
		self.addCleanup(frappe.set_user, "Administrator")
		for name in frappe.get_all(DOCTYPE, filters={"customer": ["in", [shop.CUSTOMER, shop.OTHER_CUSTOMER]]}, pluck="name"):
			frappe.delete_doc(DOCTYPE, name, force=True, ignore_permissions=True)
		for user in (shop.USER, shop.OTHER_USER):
			frappe.set_user(user)
			answer = brand_api.save_my_brand(
				{"brand_name": f"Brand of {user}", "default_markup_percent": 30, "price_rounding": "none", "price_display": "Excluding VAT"}
			)
			self.assertTrue(answer["ok"], answer)
		frappe.db.commit()
		frappe.set_user(shop.USER)

	def ask(self, fmt="xlsx", chromium=False, **kwargs):
		"""create_catalogue as the signed-in reseller, the job's enqueueing caught. Chromium is off unless asked:
		a file does not need it."""
		args = {"title": "Tariff 2027", "picked": json.dumps(PICKED), "options": json.dumps({"prices": "tariff"}), "format": fmt}
		args.update(kwargs)
		with mock.patch("frappe.enqueue") as enqueue, mock.patch.object(catalogue_api, "chromium_available", return_value=chromium):
			answer = catalogue_api.create_catalogue(**args)
		self.enqueue = enqueue
		return answer

	def newest(self, customer=shop.CUSTOMER):
		return frappe.get_all(DOCTYPE, filters={"customer": customer}, fields=["name"], order_by="creation desc", limit=1)[0].name

	def made(self, fmt="xlsx", **kwargs):
		"""A file asked for and made by the job, run inline. Returns the download."""
		answer = self.ask(fmt, **kwargs)
		self.assertTrue(answer["ok"], answer)
		name = self.newest()
		frappe.db.commit()
		jobs.generate(name)
		self.assertEqual(frappe.db.get_value(DOCTYPE, name, "status"), READY, frappe.db.get_value(DOCTYPE, name, "error"))
		return catalogue_api.download_catalogue(name=name)

	# --- asking ----------------------------------------------------------------------------------

	def test_the_format_is_kept_and_a_record_without_one_is_a_pdf(self):
		for fmt, stored in (("xlsx", "XLSX"), ("csv", "CSV")):
			with self.subTest(fmt=fmt):
				for name in frappe.get_all(DOCTYPE, filters={"customer": shop.CUSTOMER}, pluck="name"):
					frappe.delete_doc(DOCTYPE, name, force=True, ignore_permissions=True)
				self.assertTrue(self.ask(fmt)["ok"])
				self.assertEqual(frappe.db.get_value(DOCTYPE, self.newest(), "format"), stored)
		row = catalogue_api._row(frappe._dict(name="x", status="Ready", options="{}", format=None))
		self.assertEqual(row["format"], "PDF")

	def test_a_format_that_is_not_in_the_list_is_refused_naming_its_field(self):
		for value in ("docx", "xlsx; drop", "<b>", 3):
			with self.subTest(value=value):
				answer = self.ask(value)
				self.assertEqual((answer["ok"], answer["field"]), (False, "format"))

	def test_a_file_needs_no_chromium_but_a_pdf_does(self):
		self.assertTrue(self.ask("xlsx", chromium=False)["ok"])
		for name in frappe.get_all(DOCTYPE, filters={"customer": shop.CUSTOMER}, pluck="name"):
			frappe.delete_doc(DOCTYPE, name, force=True, ignore_permissions=True)
		answer = self.ask("pdf", chromium=False, options=json.dumps({"layout": "grid"}))
		self.assertFalse(answer["ok"])

	def test_the_price_modes_of_a_file_are_refused_to_a_pdf_which_is_for_their_customers(self):
		for mode in ("tariff", "both"):
			with self.subTest(mode=mode):
				answer = self.ask("pdf", chromium=True, options=json.dumps({"prices": mode}))
				self.assertEqual((answer["ok"], answer["field"]), (False, "prices"))
		self.assertFalse(frappe.get_all(DOCTYPE, filters={"customer": shop.CUSTOMER}))

	def test_a_file_with_their_own_price_needs_no_brand_and_one_with_a_selling_price_does(self):
		frappe.set_user("Administrator")
		frappe.delete_doc("Reseller Catalogue Profile", shop.OTHER_CUSTOMER, force=True, ignore_permissions=True)
		frappe.set_user(shop.OTHER_USER)
		self.assertTrue(self.ask("csv", options=json.dumps({"prices": "tariff"}))["ok"])
		for name in frappe.get_all(DOCTYPE, filters={"customer": shop.OTHER_CUSTOMER}, pluck="name"):
			frappe.delete_doc(DOCTYPE, name, force=True, ignore_permissions=True)
		for mode in ("customers", "both"):
			with self.subTest(mode=mode):
				answer = self.ask("csv", options=json.dumps({"prices": mode}))
				self.assertEqual((answer["ok"], answer["message"]), (False, catalogue_message("no_brand_prices")))

	def test_a_file_over_the_cap_is_refused_at_once_in_lines(self):
		with mock.patch.object(rows, "MAX_ROWS", 2):
			answer = self.ask("csv", picked=json.dumps({"items": [shop.MODEL, shop.FIRST]}))
		self.assertEqual((answer["ok"], answer["field"]), (False, "selection"))
		self.assertEqual(answer["message"], catalogue_message("too_many_rows", 2))
		self.assertFalse(frappe.get_all(DOCTYPE, filters={"customer": shop.CUSTOMER}))

	def test_the_count_of_a_file_is_in_lines_and_a_models_variants_count_each(self):
		picked = json.dumps({"items": [shop.MODEL, shop.FIRST]})
		file_count = catalogue_api.count_products(picked=picked, format="csv")
		pdf_count = catalogue_api.count_products(picked=picked)
		self.assertEqual((file_count["count"], file_count["unit"], file_count["limit"]), (3, "lines", rows.MAX_ROWS))
		self.assertEqual((pdf_count["count"], pdf_count["unit"]), (2, "products"))

	def test_every_endpoint_still_refuses_a_visitor_an_unallowed_account_and_staff(self):
		for user in ("Guest", shop.OUT_USER, "Administrator"):
			frappe.set_user(user)
			with self.subTest(user=user):
				with self.assertRaises(frappe.PermissionError):
					catalogue_api.create_catalogue(title="x", picked=json.dumps(PICKED), format="csv")
				with self.assertRaises(frappe.PermissionError):
					catalogue_api.count_products(picked=json.dumps(PICKED), format="csv")

	# --- making and downloading ----------------------------------------------------------------------

	def test_a_workbook_is_made_and_downloaded_as_an_excel_file(self):
		response = self.made("xlsx", options=json.dumps({"prices": "both"}))
		self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
		self.assertIn("Tariff-2027.xlsx", response.headers["Content-Disposition"])
		self.assertEqual(response.headers["X-Robots-Tag"], "noindex")
		self.assertEqual(response.headers["Cache-Control"], "private, no-store")
		workbook = load_workbook(io.BytesIO(response.get_data()))
		sheet = workbook.worksheets[0]
		header = [cell.value for cell in sheet[1]]
		self.assertEqual(len(header), len(rows.columns_for("both")))
		codes = [sheet.cell(row=index, column=1).value for index in range(2, sheet.max_row + 1)]
		self.assertIn(shop.FIRST, codes)
		self.assertEqual(len(workbook.worksheets), 2)
		record = frappe.get_doc(DOCTYPE, self.newest())
		self.assertEqual((record.item_count, record.page_count), (len(codes), 0))

	def test_a_csv_file_is_made_and_downloaded_as_a_csv_file(self):
		response = self.made("csv", options=json.dumps({"prices": "both"}))
		self.assertEqual(response.mimetype, "text/csv")
		self.assertIn("Tariff-2027.csv", response.headers["Content-Disposition"])
		records = list(csv.DictReader(io.StringIO(response.get_data(as_text=True), newline="")))
		by_code = {record["code"]: record for record in records}
		self.assertEqual((by_code[shop.FIRST]["purchase_price"], by_code[shop.FIRST]["price"]), ("80.00", "104.00"))
		self.assertNotIn(shop.HIDDEN, by_code)

	def test_the_mode_for_their_customers_carries_no_purchase_price_in_a_file_either(self):
		response = self.made("csv", options=json.dumps({"prices": "customers"}))
		text = response.get_data(as_text=True)
		self.assertNotIn("purchase_price", text)
		record = {record["code"]: record for record in csv.DictReader(io.StringIO(text, newline=""))}[shop.FIRST]
		self.assertEqual(record["price"], "104.00")

	def test_making_it_again_keeps_the_format_and_the_choices(self):
		self.made("csv", options=json.dumps({"prices": "tariff", "markup_percent": 10}))
		first = self.newest()
		with mock.patch("frappe.enqueue"), mock.patch.object(catalogue_api, "chromium_available", return_value=False):
			answer = catalogue_api.remake_catalogue(name=first)
		self.assertTrue(answer["ok"], answer)
		newest = frappe.get_doc(DOCTYPE, self.newest())
		self.assertNotEqual(newest.name, first)
		self.assertEqual(newest.format, "CSV")
		self.assertEqual(json.loads(newest.options), {"markup_percent": 10.0, "prices": "tariff"})

	def test_a_file_the_job_cannot_make_says_so_in_a_sentence(self):
		answer = self.ask("csv")
		self.assertTrue(answer["ok"])
		name = self.newest()
		frappe.db.commit()
		with mock.patch.object(rows, "MAX_ROWS", 1), mock.patch.dict(jobs.NUMBERS, {"too_many_rows": 1}):
			jobs.generate(name)
		record = frappe.get_doc(DOCTYPE, name)
		self.assertEqual(record.status, FAILED)
		self.assertEqual(record.error, catalogue_message("too_many_rows", 1))
		with self.assertRaises(frappe.DoesNotExistError):
			catalogue_api.download_catalogue(name=name)

	def test_the_second_reseller_cannot_download_the_first_ones_file(self):
		self.made("csv")
		name = self.newest()
		frappe.set_user(shop.OTHER_USER)
		with self.assertRaises(frappe.DoesNotExistError):
			catalogue_api.download_catalogue(name=name)
		self.assertEqual(catalogue_api.get_catalogues()["catalogues"], [])
