# //// Neoffice — added file (no upstream equivalent).
"""The three files a table of products becomes: CSV and JSON for another system, Excel for a person.

Built from lines written by hand, so that nothing here depends on a shop: what is proved is the shape of the
files — the same keys in the same order whatever the language, figures that stay figures, a cell that starts
like a formula that stays text in Excel — and that a file never says more than its lines do.
"""

import csv
import io
import json
import unittest

from openpyxl import load_workbook

from webshop.webshop.catalogue import rows, tables


def line(**fields):
	base = {
		"code": "SKU-1",
		"model": "",
		"name": "Trail shoe",
		"brand": "Alpine",
		"category": "Shoes > Trail",
		"ean": "4006381333931",
		"mpn": "AB-12",
		"unit": "Nos",
		"attributes": {},
		"currency": "CHF",
		"price": 104.0,
		"price_includes_vat": False,
		"availability": "in_stock",
		"description": "Light.\nSecond line, with a comma.",
		"image": "https://shop.example/files/a.jpg",
		"more_images": ["https://shop.example/files/b.jpg", "https://shop.example/files/c.jpg"],
		"url": "https://shop.example/trail-shoe",
	}
	base.update(fields)
	return base


def data(lines=None, prices="customers"):
	import frappe

	return frappe._dict(
		columns=rows.columns_for(prices),
		rows=lines if lines is not None else [line(), line(code="SKU-2", model="SKU-0", attributes={"Size": "42", "Colour": "Red"}, price=62.4)],
		currency="CHF",
		prices=prices,
		markup_percent=30.0,
		including_vat=False,
		prices_as_of="2026-10-31",
		title="Spring 2027",
		issuer="Alpine Distribution SA",
	)


class TestCsv(unittest.TestCase):
	def read(self, content):
		return list(csv.reader(io.StringIO(content.decode("utf-8"), newline="")))

	def test_the_header_is_the_stable_machine_names_in_the_column_order(self):
		table = self.read(tables.csv_bytes(data()))
		self.assertEqual(table[0], rows.columns_for("customers"))
		self.assertEqual(len(table), 3)
		self.assertTrue(all(len(record) == len(table[0]) for record in table))

	def test_it_is_utf_8_without_a_byte_order_mark_and_lines_end_with_crlf(self):
		content = tables.csv_bytes(data([line(name="Chaussure d'été à lacets")]))
		self.assertFalse(content.startswith(b"\xef\xbb\xbf"))
		self.assertIn(b"\r\n", content)
		self.assertIn("Chaussure d'été à lacets", content.decode("utf-8"))

	def test_a_comma_a_quote_and_a_line_break_in_a_cell_come_back_as_they_went_in(self):
		record = self.read(tables.csv_bytes(data([line(name='The "best", really', description="a\nb, c")])))[1]
		columns = rows.columns_for("customers")
		self.assertEqual(record[columns.index("name")], 'The "best", really')
		self.assertEqual(record[columns.index("description")], "a\nb, c")

	def test_figures_have_a_dot_and_two_decimals_and_a_missing_one_is_an_empty_cell(self):
		columns = rows.columns_for("customers")
		records = self.read(tables.csv_bytes(data([line(price=104.0), line(code="SKU-9", price=None)])))
		self.assertEqual(records[1][columns.index("price")], "104.00")
		self.assertEqual(records[2][columns.index("price")], "")

	def test_a_price_keeps_the_decimals_it_has(self):
		columns = rows.columns_for("tariff")
		records = self.read(tables.csv_bytes(data([line(purchase_price=0.125), line(purchase_price=80.0)], "tariff")))
		self.assertEqual(records[1][columns.index("purchase_price")], "0.125")
		self.assertEqual(records[2][columns.index("purchase_price")], "80.00")

	def test_lists_and_flags_are_written_flat(self):
		columns = rows.columns_for("customers")
		record = self.read(tables.csv_bytes(data([line(attributes={"Size": "42", "Colour": "Red"})])))[1]
		self.assertEqual(record[columns.index("attributes")], "Size: 42; Colour: Red")
		self.assertEqual(record[columns.index("more_images")], "https://shop.example/files/b.jpg,https://shop.example/files/c.jpg")
		self.assertEqual(record[columns.index("price_includes_vat")], "false")

	def test_a_text_that_looks_like_a_formula_is_not_altered_in_a_file_for_another_system(self):
		record = self.read(tables.csv_bytes(data([line(name="=1+1", code="-12")])))[1]
		columns = rows.columns_for("customers")
		self.assertEqual((record[columns.index("name")], record[columns.index("code")]), ("=1+1", "-12"))

	def test_only_the_columns_of_the_mode_are_written(self):
		self.assertNotIn("price", self.read(tables.csv_bytes(data(prices="none")))[0])


class TestJson(unittest.TestCase):
	def test_a_document_with_the_figures_of_the_file_and_one_object_per_line(self):
		document = json.loads(tables.json_bytes(data(), generated_on="2026-10-31T08:00:00"))
		self.assertEqual(document["generated_on"], "2026-10-31T08:00:00")
		self.assertEqual((document["currency"], document["prices_as_of"], document["count"]), ("CHF", "2026-10-31", 2))
		first, second = document["products"]
		self.assertEqual(first["code"], "SKU-1")
		self.assertEqual(list(first), rows.columns_for("customers"), "the keys are the CSV's columns, in its order")
		self.assertEqual(second["attributes"], {"Size": "42", "Colour": "Red"})
		self.assertEqual(first["more_images"], ["https://shop.example/files/b.jpg", "https://shop.example/files/c.jpg"])

	def test_prices_are_numbers_and_a_missing_one_is_null(self):
		first, second = json.loads(tables.json_bytes(data([line(price=104.0), line(price=None)])))["products"]
		self.assertEqual(first["price"], 104.0)
		self.assertIsInstance(first["price"], float)
		self.assertIsNone(second["price"])
		self.assertIs(first["price_includes_vat"], False)

	def test_accents_are_written_as_they_are(self):
		self.assertIn("à lacets".encode(), tables.json_bytes(data([line(name="à lacets")])))

	def test_the_document_never_names_a_price_the_mode_left_out(self):
		text = tables.json_bytes(data(prices="none")).decode()
		self.assertNotIn('"price"', text)
		self.assertNotIn("purchase_price", text)


class TestExcel(unittest.TestCase):
	LABELS = {
		"columns": {key: key.replace("_", " ").title() for key in rows.columns_for("both")},
		"availability": {"in_stock": "In stock", "out_of_stock": "Out of stock", "backorder": "On backorder"},
		"yes": "Yes",
		"no": "No",
		"information": "Information",
		"products": "Products",
		"facts": {"Title": "Spring 2027"},
	}

	def book(self, **kwargs):
		return load_workbook(io.BytesIO(tables.xlsx_bytes(kwargs.pop("table", None) or data(), self.LABELS)))

	def test_a_sheet_of_products_and_a_sheet_of_information(self):
		workbook = self.book()
		self.assertEqual(workbook.sheetnames, ["Products", "Information"])

	def test_the_header_is_the_readable_names_bold_and_frozen_with_a_filter(self):
		sheet = self.book()["Products"]
		header = [cell.value for cell in sheet[1]]
		self.assertEqual(header[:3], ["Code", "Model", "Name"])
		self.assertTrue(all(cell.font.bold for cell in sheet[1]))
		self.assertEqual(sheet.freeze_panes, "A2")
		self.assertTrue(sheet.auto_filter.ref)

	def test_a_price_is_a_number_with_two_decimals_not_a_text(self):
		sheet = self.book()["Products"]
		column = rows.columns_for("customers").index("price") + 1
		cell = sheet.cell(row=2, column=column)
		self.assertEqual(cell.value, 104.0)
		self.assertEqual(cell.data_type, "n")
		self.assertIn("0.00", cell.number_format)

	def test_a_cell_that_starts_like_a_formula_stays_text(self):
		sheet = self.book(table=data([line(name="=HYPERLINK(\"http://x\",\"y\")", description="+1 deal", brand="@SUM(A1)")]))["Products"]
		columns = rows.columns_for("customers")
		for key, value in (("name", '=HYPERLINK("http://x","y")'), ("description", "+1 deal"), ("brand", "@SUM(A1)")):
			cell = sheet.cell(row=2, column=columns.index(key) + 1)
			with self.subTest(key=key):
				self.assertEqual((cell.value, cell.data_type), (value, "s"))

	def test_availability_and_flags_are_worded_for_a_person(self):
		sheet = self.book()["Products"]
		columns = rows.columns_for("customers")
		self.assertEqual(sheet.cell(row=2, column=columns.index("availability") + 1).value, "In stock")
		self.assertEqual(sheet.cell(row=2, column=columns.index("price_includes_vat") + 1).value, "No")

	def test_a_list_of_pictures_is_one_cell(self):
		sheet = self.book()["Products"]
		value = sheet.cell(row=2, column=rows.columns_for("customers").index("more_images") + 1).value
		self.assertEqual(value, "https://shop.example/files/b.jpg\nhttps://shop.example/files/c.jpg")

	def test_a_control_character_in_a_text_does_not_stop_the_workbook(self):
		"""A vertical tab pasted from a word processor made openpyxl refuse the whole file (IllegalCharacterError)."""
		labels = {**self.LABELS, "facts": {"Title": "Spring\x002027"}}
		table = data([line(name="Trail\x0bshoe", brand="Al\x01pine", description="two\x08 words", more_images=["https://shop.example/a\x0c.jpg"], attributes={"Siz\x02e": "4\x1f2"})])
		sheet = load_workbook(io.BytesIO(tables.xlsx_bytes(table, labels)))
		columns = rows.columns_for("customers")
		cells = {key: sheet["Products"].cell(row=2, column=columns.index(key) + 1).value for key in ("name", "brand", "description", "attributes")}
		self.assertEqual(cells, {"name": "Trailshoe", "brand": "Alpine", "description": "two words", "attributes": "Size: 42"})
		self.assertEqual(sheet["Information"].cell(row=1, column=2).value, "Spring2027")
		self.assertIn("Trail\x0bshoe".encode(), tables.csv_bytes(table), "the CSV is left as it is")

	def test_the_information_sheet_says_what_the_file_is(self):
		sheet = self.book()["Information"]
		cells = {row[0].value: row[1].value for row in sheet.iter_rows() if row[0].value}
		self.assertEqual(cells["Title"], "Spring 2027")

	def test_a_file_with_no_line_is_still_a_file(self):
		sheet = self.book(table=data([]))["Products"]
		self.assertEqual(sheet.max_row, 1)

	def test_only_the_columns_of_the_mode_are_in_the_sheet(self):
		sheet = self.book(table=data(prices="none"))["Products"]
		self.assertNotIn("Price", [cell.value for cell in sheet[1]])
